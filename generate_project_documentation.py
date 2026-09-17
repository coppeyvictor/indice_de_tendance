from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


OUTPUT_PATH = Path("Documentation_Projet_Fear_Greed.docx")


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_function_table(document: Document, functions: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Fonction"
    table.rows[0].cells[1].text = "Rôle"
    for name, description in functions:
        cells = table.add_row().cells
        cells[0].text = name
        cells[1].text = description


def add_module(document: Document, filename: str, purpose: str, functions: list[tuple[str, str]]) -> None:
    add_heading(document, filename, level=2)
    document.add_paragraph(purpose)
    if functions:
        add_function_table(document, functions)


def build_document() -> Path:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)

    title = document.add_heading("Documentation du projet", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Fear & Greed Dashboard - Finance, ESG et Ecologie")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("Projet Python local de collecte d'articles, analyse de sentiment et visualisation interactive.")

    add_heading(document, "1. Objectif", level=1)
    document.add_paragraph(
        "Le projet construit des indices de sentiment a partir d'articles de presse. "
        "Il propose trois vues : Finance (Fear & Greed Index), ESG (ESG Sentiment Index) "
        "et Ecologie (Ecology Sentiment Index). Les indices sont compris entre 0 et 100."
    )
    add_bullets(document, [
        "Finance : mesure du sentiment lie aux marches, a l'economie et aux crypto-actifs.",
        "ESG : mesure de la tonalite des articles lies a la finance durable, aux normes et a la transition.",
        "Ecologie : mesure de la tonalite des articles lies au climat, a la biodiversite et aux energies propres.",
        "Le tableau de bord est une application Flask locale et ne constitue pas un conseil financier.",
    ])

    add_heading(document, "2. Architecture et flux de donnees", level=1)
    document.add_paragraph(
        "1. veille_presse.py recupere les flux RSS et classe les articles. "
        "2. sentiment.py attribue un score de sentiment. "
        "3. cluster_articles.py regroupe les articles parlant du meme evenement. "
        "4. fear_greed_index.py calcule les indices et genere le HTML Plotly. "
        "5. app.py sert le tableau de bord via Flask."
    )
    add_heading(document, "Base de donnees", level=2)
    document.add_paragraph("La base SQLite sentiments.db contient notamment deux tables :")
    add_bullets(document, [
        "daily_articles : articles collectes, date, source, texte, score, cluster, categorie, pays, type d'actif et theme.",
        "articles : archive des articles analyses et de leur score de sentiment.",
    ])

    add_heading(document, "3. Calcul de l'indice", level=1)
    document.add_paragraph("Chaque score de sentiment est borne entre -1 et +1 et est converti ainsi :")
    document.add_paragraph("Indice = (score moyen + 1) x 50", style="Intense Quote")
    add_bullets(document, [
        "0-20 : peur extreme / tres preoccupant.",
        "20-40 : peur / preoccupant.",
        "40-60 : neutre / mixte.",
        "60-80 : greed / positif.",
        "80-100 : greed extreme / tres positif.",
        "Les clusters limitent le poids de plusieurs articles redondants sur un meme evenement.",
        "Les moyennes 3, 7 et 30 jours sont ponderees par le nombre d'articles disponibles.",
        "Pour une journee sans score, la vue globale peut reporter la derniere valeur connue afin d'afficher la date du jour.",
    ])

    add_heading(document, "4. Fichiers et fonctions importantes", level=1)
    add_module(document, "app.py", "Point d'entree de l'application web Flask. Genere le tableau de bord demande et le renvoie au navigateur.", [
        ("normalize_theme(theme)", "Valide le theme demande : finance, esg ou ecology."),
        ("build_dashboard(days, theme)", "Genere fear_greed_chart.html a partir des donnees SQLite."),
        ("dashboard_summary(days)", "Construit les indicateurs resumes utilises par l'API."),
        ("index() et dashboard()", "Routes Flask qui servent le tableau de bord."),
        ("api_dashboard()", "Route JSON avec le theme et le resume des donnees."),
    ])
    add_module(document, "veille_presse.py", "Collecteur RSS et couche d'ingestion des articles dans SQLite. Il contient les listes de sources Finance, ESG et Ecologie.", [
        ("live_google_news_urls()", "Construit les recherches Google News recentes dans plusieurs langues."),
        ("classify_theme(...)", "Classe un article en finance, esg ou ecology selon ses mots-cles."),
        ("classify_article(...)", "Classe un article en economie, marches ou crypto."),
        ("classify_country(...) / classify_asset_type(...)", "Ajoute le pays et le type d'actif ou de sujet."),
        ("initialize_database()", "Cree ou migre les tables SQLite et recalcule les metadonnees."),
        ("get_articles(...) / collect_daily_articles(...)", "Telecharge, dedoublonne, traduit les titres et enregistre les articles RSS."),
        ("theme_rss_urls(theme)", "Retourne les flux specifiques a un theme, y compris les sources ESG et climat."),
        ("process_new_articles(...) ", "Collecte les articles non notes, analyse leur sentiment puis sauvegarde les scores."),
        ("force_reanalyze_theme(...) ", "Remet les scores d'un theme a null pour forcer une nouvelle analyse."),
    ])
    add_module(document, "sentiment.py", "Moteur d'analyse de sentiment. Le mode local utilise FinBERT ; Gemini peut etre active par variable d'environnement.", [
        ("parse_score(response_text)", "Valide un score texte strictement compris entre -1 et +1."),
        ("parse_sentiment_analysis(response_text)", "Valide la reponse JSON Gemini et combine sentiment, impact et confiance."),
        ("prepare_sentiment_chunks(text)", "Decoupe les textes longs sans perdre leur contenu."),
        ("analyze_sentiment_local_details(text)", "Analyse locale FinBERT et retourne score, impact et confiance."),
        ("analyze_sentiment(text)", "Choisit le fournisseur local ou Gemini selon SENTIMENT_PROVIDER."),
    ])
    add_module(document, "cluster_articles.py", "Regroupe les articles similaires pour limiter l'influence des doublons dans l'indice.", [
        ("tokenize(text)", "Nettoie et tokenise un texte en supprimant les mots vides."),
        ("tfidf_vectors(documents)", "Construit les vecteurs TF-IDF des articles."),
        ("cosine_similarity(left, right)", "Calcule la similarite entre deux vecteurs."),
        ("cluster_documents(documents, threshold)", "Attribue un identifiant de cluster aux documents similaires."),
        ("cluster_period(start_date, end_date, threshold)", "Ecrit les clusters et leurs tailles dans daily_articles."),
    ])
    add_module(document, "fear_greed_index.py", "Coeur de calcul des indices et generateur de la page HTML interactive Plotly.", [
        ("compute_fear_greed_index(score)", "Convertit un score -1 a +1 vers une echelle 0 a 100."),
        ("_cluster_rows(...) / _aggregate_clusters(...)", "Lit les articles notes et calcule les moyennes ponderees par cluster."),
        ("get_daily_sentiment_index(...) ", "Produit une serie quotidienne globale ou par categorie, avec report de valeur optionnel."),
        ("get_weighted_fear_greed_index(days)", "Calcule l'indice global pondere par le volume d'articles."),
        ("_build_windowed_curve(...) / get_index_curves(...)", "Construit les courbes Daily, 3, 7 et 30 jours."),
        ("generate_interactive_fear_greed_chart(...)", "Genere le dashboard HTML, les controles de theme, de periode, les filtres d'articles et le graphique Plotly."),
    ])
    add_module(document, "collect_history.py", "Collecte et analyse une plage historique de dates pour augmenter la profondeur des donnees.", [
        ("iter_days(start_date, end_date)", "Itere sur toutes les dates d'une periode."),
        ("collect_history(...) ", "Collecte les flux recents puis les recherches historiques quand le volume journalier est faible."),
    ])
    add_module(document, "dashboard.py", "Affiche un resume textuel de l'indice dans le terminal.", [
        ("fetch_daily_rows(days)", "Prepare les lignes quotidiennes utilisables."),
        ("dashboard_summary(days)", "Calcule score moyen, indice, statut et detail journalier."),
        ("print_dashboard(days)", "Imprime le resume dans le terminal."),
    ])
    add_module(document, "open_dashboard.py", "Genere fear_greed_chart.html et l'ouvre eventuellement dans le navigateur.", [
        ("main()", "Lit les options --days, --theme et --no-open puis genere le fichier HTML."),
    ])
    add_module(document, "indice_sentiment.py", "Script de demonstration qui mesure le temps d'analyse d'un texte exemple.", [("main()", "Analyse un texte financier et affiche le score et la latence.")])
    add_module(document, "lien API.py", "Test de connectivite minimal vers le fournisseur de sentiment.", [("main()", "Analyse une phrase courte et affiche le score obtenu.")])
    add_module(document, "quota_gemini.py", "Verifie si une requete Gemini est acceptee et extrait les informations de quota en cas de limite.", [("extract_detail(...)", "Extrait une information dans le message d'erreur."), ("main()", "Execute une requete minimale et interprete une erreur de quota.")])
    add_module(document, "test_gemini.py", "Suite pytest couvrant les calculs, classifications, clustering et validateurs importants.", [])
    add_module(document, "templates/index.html", "Ancien template Flask de presentation. La route principale sert actuellement le HTML genere par fear_greed_index.py.", [])

    add_heading(document, "5. Utilisation", level=1)
    add_heading(document, "Installation", level=2)
    document.add_paragraph("python3 -m venv .venv\nsource .venv/bin/activate\npip install -r requirements.txt", style="Intense Quote")
    add_heading(document, "Lancer l'application Flask", level=2)
    document.add_paragraph("PORT=5003 python3 app.py\nPuis ouvrir : http://127.0.0.1:5003", style="Intense Quote")
    add_heading(document, "Collecter les articles du jour", level=2)
    document.add_paragraph("python3 veille_presse.py --date 2026-09-15 --theme finance --once\npython3 veille_presse.py --date 2026-09-15 --theme esg --once\npython3 veille_presse.py --date 2026-09-15 --theme ecology --once", style="Intense Quote")
    add_heading(document, "Verifier les volumes par theme", level=2)
    document.add_paragraph("sqlite3 -header -column sentiments.db \"SELECT theme, COUNT(*) AS articles FROM daily_articles GROUP BY theme ORDER BY articles DESC;\"", style="Intense Quote")
    add_heading(document, "Executer les tests", level=2)
    document.add_paragraph("python3 -m pytest test_gemini.py", style="Intense Quote")

    add_heading(document, "6. Limites et precautions", level=1)
    add_bullets(document, [
        "Les flux RSS peuvent etre indisponibles, ralentis ou modifier leur format.",
        "Le classement par mots-cles et les scores automatiques peuvent contenir des erreurs.",
        "Le volume d'articles est inegal selon les themes et les dates.",
        "Les resultats sont destines a l'etude et a la visualisation, pas a la prise de decision financiere autonome.",
        "Une cle Gemini doit etre conservee dans la variable d'environnement GEMINI_API_KEY, jamais dans le code.",
    ])

    document.save(OUTPUT_PATH)
    return OUTPUT_PATH.resolve()


if __name__ == "__main__":
    print(build_document())